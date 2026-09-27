"""Resume exports share one ordered block model with the in-app preview.

No generation or embellishment happens here. Only supplied, confirmed resume fields
are typeset. A4, a single column, ordinary paragraphs, and black text keep the DOCX
stable across Word and WPS; PDF embeds a Chinese font whenever one is available.
"""
from __future__ import annotations

import os
from pathlib import Path
from typing import Literal
from xml.sax.saxutils import escape

from pydantic import BaseModel, Field


class ResumeProject(BaseModel):
    id: str = ""
    title: str = ""
    context: str = ""
    bullets: list[str] = Field(default_factory=list)


class ResumeContent(BaseModel):
    name: str = ""
    headline: str = ""
    contact: str = ""
    summary: str = ""
    skills: list[str] = Field(default_factory=list)
    projects: list[ResumeProject] = Field(default_factory=list)
    education: list[str] = Field(default_factory=list)
    experience: list[str] = Field(default_factory=list)
    language: Literal["zh", "en"] = "zh"


def resume_blocks(content: dict) -> list[dict[str, str]]:
    """A deterministic source for preview and both file formats; preserves wording."""
    resume = ResumeContent.model_validate(content)
    if not resume.experience and not any(p.context or p.bullets for p in resume.projects):
        raise ValueError("个人经历不足：请先确认至少一段工作经历或含具体内容的项目经历。")
    labels = ({"summary": "个人概述", "skills": "技能", "experience": "工作经历", "projects": "项目经历", "education": "教育经历"}
              if resume.language == "zh" else
              {"summary": "Summary", "skills": "Skills", "experience": "Experience", "projects": "Projects", "education": "Education"})
    blocks: list[dict[str, str]] = []

    def add(kind: str, text: str):
        if text and text.strip():
            blocks.append({"kind": kind, "text": text.strip()})

    add("title", resume.name)
    add("subtitle", resume.headline)
    add("contact", resume.contact)
    if resume.summary:
        add("heading", labels["summary"])
        add("body", resume.summary)
    for field in ("skills", "experience"):
        values = getattr(resume, field)
        if any(value.strip() for value in values):
            add("heading", labels[field])
            for value in values:
                add("bullet", value)
    if resume.projects:
        add("heading", labels["projects"])
        for project in resume.projects:
            add("subheading", project.title)
            add("body", project.context)
            for value in project.bullets:
                add("bullet", value)
    if any(value.strip() for value in resume.education):
        add("heading", labels["education"])
        for value in resume.education:
            add("body", value)
    return blocks


def _export_docx(blocks: list[dict], output: Path) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = next((b["text"] for b in blocks if b["kind"] == "title"), "")
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.0)
    # Override Word theme colours and both Western/East Asian font slots: WPS
    # otherwise sometimes inherits a blue heading colour or an inconsistent font.
    for style in document.styles:
        if hasattr(style, "font"):
            style.font.name = "Arial"
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.size = Pt(10.5)
            rpr = style.element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                rpr.insert(0, fonts)
            for key in ("ascii", "hAnsi", "cs"):
                fonts.set(qn("w:" + key), "Arial")
            fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                fonts.attrib.pop(qn("w:" + key), None)
            color = rpr.find(qn("w:color"))
            if color is not None:
                for key in ("themeColor", "themeTint", "themeShade"):
                    color.attrib.pop(qn("w:" + key), None)
    for block in blocks:
        kind, text = block["kind"], block["text"]
        style = {"title": "Title", "heading": "Heading 1", "subheading": "Heading 2"}.get(kind, "Normal")
        paragraph = document.add_paragraph(style=style)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.2
        paragraph.paragraph_format.widow_control = True
        if kind in {"heading", "subheading", "title", "subtitle"}:
            paragraph.paragraph_format.keep_with_next = True
        if kind == "heading":
            paragraph.paragraph_format.space_before = Pt(12)
        if kind == "bullet":
            paragraph.paragraph_format.left_indent = Cm(0.35)
            paragraph.paragraph_format.first_line_indent = Cm(-0.3)
            text = "• " + text
        run = paragraph.add_run(text)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.size = Pt({"title": 22, "subtitle": 12, "heading": 13, "subheading": 11}.get(kind, 10.5))
        run.bold = kind in {"title", "heading", "subheading"}
    document.save(output)


def _pdf_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_path = os.getenv("RESUME_FONT_PATH", "")
    candidates = [font_path, "C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simhei.ttf",
                  "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
                  "/usr/share/fonts/truetype/arphic/uming.ttc"]
    for candidate in candidates:
        if not candidate or not Path(candidate).is_file():
            continue
        try:
            font = "ResumeEmbedded"
            if font not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font, candidate, subfontIndex=0))
            return font
        except Exception:
            continue
    # Standard CJK CID font remains selectable/searchable; explicitly documented
    # as fallback. Set RESUME_FONT_PATH to embed a portable TrueType Chinese font.
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    if "STSong-Light" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def _export_pdf(blocks: list[dict], output: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    font = _pdf_font()
    styles = {}
    for kind, size in {"title": 22, "subtitle": 12, "contact": 10.5, "heading": 13,
                       "subheading": 11, "body": 10.5, "bullet": 10.5}.items():
        styles[kind] = ParagraphStyle(kind, fontName=font, fontSize=size, leading=size * 1.35,
            textColor=colors.black, alignment=TA_LEFT, spaceAfter=5,
            spaceBefore=12 if kind == "heading" else 0,
            keepWithNext=kind in {"title", "subtitle", "heading", "subheading"},
            wordWrap="CJK", allowWidows=0, allowOrphans=0,
            leftIndent=10 if kind == "bullet" else 0, firstLineIndent=-8 if kind == "bullet" else 0)
    story = []
    for block in blocks:
        kind = block["kind"]
        value = escape(block["text"]).replace("\n", "<br/>")
        if kind == "bullet":
            value = "• " + value
        story.append(Paragraph(value, styles[kind]))
    document = SimpleDocTemplate(str(output), pagesize=A4, rightMargin=2 * cm, leftMargin=2 * cm,
        topMargin=1.8 * cm, bottomMargin=1.8 * cm, title="", author="", allowSplitting=1)
    document.build(story)


def export_resume(content: dict, fmt: str, output: Path) -> Path:
    """Export the supplied version only. The caller owns base/version persistence."""
    fmt = fmt.lower().lstrip(".")
    if fmt not in {"docx", "pdf"}:
        raise ValueError("仅支持 DOCX 和 PDF 导出。")
    blocks = resume_blocks(content)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "docx":
        _export_docx(blocks, output)
    else:
        _export_pdf(blocks, output)
    return output
